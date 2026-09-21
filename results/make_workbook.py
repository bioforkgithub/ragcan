#!/usr/bin/env python3
"""Build RaGCAn_results.xlsx: one sheet per organism group, one row per database genus per screened set.

This is the record of how RaGCAn_results.xlsx in this folder was built. It reads the survey and
extension working folders (the per-genus proposed_bins.tsv files and the GTDB taxonomy), which are
not in this repository because of their size; rerunning the survey recreates them.

Columns: DB_searched | DB_agreement | Genus | Screened set | RaGCAn_1, RaGCAn_2, ...
  DB_searched   GTDB (R232) for prokaryote genomes GTDB classifies; NCBI for eukaryotes and for
                prokaryote genomes absent from GTDB (genus = NCBI organism name).
  DB_agreement  yes when all of this genus's genomes in the set fall in one RaGCAn bin and that bin
                holds no genome of any other genus; otherwise no.
  RaGCAn_k      accessions of this genus's genomes in bin k of that set (bins as numbered by the program,
                65% AAI).
Every value is read from the program's proposed_bins.tsv files and the genome manifests.
"""
import csv, re, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Results_Workbook' / 'RaGCAn_results.xlsx'
ACC = re.compile(r'(GC[AF]_\d+\.\d+)')

def tsv(p):
    return list(csv.DictReader(open(p, newline=''), delimiter='\t'))

# GTDB R232 genus per accession
gtdb = {}
for f in ('bac120_taxonomy.tsv', 'ar53_taxonomy.tsv'):
    for line in open(ROOT / 'Genus_Extension' / 'reference' / f):
        a, t = line.rstrip('\n').split('\t', 1)
        if a[:3] in ('RS_', 'GB_'): a = a.split('_', 1)[1]
        m = re.search(r'g__([^;]*)', t)
        gtdb[a] = m.group(1) if m else ''

# genome manifests: accession -> (name, ncbi genus, ftp)
info = {}
def add_selection(p):
    for r in tsv(p):
        info[r['accession']] = (r['organism_name'], r['genus'], r['ftp_path'].rstrip('/'))
add_selection(ROOT / 'LPSN_Genus_Survey' / 'reports' / 'genome_selection.tsv')
add_selection(ROOT / 'Genus_Extension' / 'reports' / 'genome_selection.tsv')

s17 = [l.rstrip('\n').split('\t') for l in open(ROOT / 'nmeth1' / 'Supplementary' / 'File_S17_eukaryote_sets.tsv')]
euk_genus = {}
start = False
for f in s17:
    if f[:3] == ['set', 'organism', 'accession']: start = True; continue
    if start:
        if len(f) < 5 or f[0].startswith('#'): break
        euk_genus[f[2]] = f[3]
        info.setdefault(f[2], (f[1], f[3], ''))
# eukaryote ftp paths from the assembly summaries used by prepare.py
for p in (ROOT / 'Eukaryote_Screen' / 'reference').glob('as_*.txt'):
    for line in open(p):
        if line.startswith('#'): continue
        x = line.rstrip('\n').split('\t')
        if x[0] in euk_genus and len(x) > 19:
            n, g, _ = info[x[0]]
            info[x[0]] = (n, g, x[19].replace('ftp://', 'https://').rstrip('/'))

def read_bins(pb):
    bins = []
    for row in tsv(pb):
        bins.append([m.group(1) for x in row['members'].split() if (m := ACC.search(x))])
    return bins

def rows_for_set(set_name, bins, eukaryote):
    lab = {}
    for b in bins:
        for a in b:
            if eukaryote: lab[a] = ('NCBI', euk_genus[a])
            elif gtdb.get(a): lab[a] = ('GTDB', gtdb[a])
            else: lab[a] = ('NCBI', info[a][1])
    bin_genera = [{lab[a][1] for a in b} for b in bins]
    keys = sorted({lab[a] for b in bins for a in b}, key=lambda k: (k[0] != 'GTDB', k[1].lower()))
    out = []
    for k in keys:
        cells = [[a for a in b if lab[a] == k] for b in bins]
        occupied = [i for i, c in enumerate(cells) if c]
        agree = len(occupied) == 1 and bin_genera[occupied[0]] == {k[1]}
        out.append([k[0], 'yes' if agree else 'no', k[1], set_name] + [', '.join(c) for c in cells])
    return out

groups = collections.OrderedDict()
def add(sheet, set_name, pb, eukaryote=False):
    groups.setdefault(sheet, []).extend(rows_for_set(set_name, read_bins(pb), eukaryote))

ext = {r['genus']: r['domain'] for r in tsv(ROOT / 'Genus_Extension' / 'phase3' / 'extension_per_genus.tsv')}
surv_arch = set()
for l in open(ROOT / 'GENERA_TESTED.txt'):
    f = l.split()
    if len(f) >= 7 and 'Archaea' in f and f[0] not in ext: surv_arch.add(f[0])

used = set()
for base, tag in (('LPSN_Genus_Survey', '4+'), ('Genus_Extension', '2-3')):
    for d in sorted((ROOT / base / 'results').iterdir()):
        pb = d / 'reports' / 'proposed_bins.tsv'
        if not pb.is_file(): continue
        arch = (ext.get(d.name) == 'Archaea') if tag == '2-3' else (d.name in surv_arch)
        sheet = f"{'Archaea' if arch else 'Bacteria'} ({tag} species)"
        add(sheet, d.name, pb)
add('Enterobacteriaceae (family)', 'Enterobacteriaceae', ROOT / 'Enterobacterales_Family' / 'run' / 'reports' / 'proposed_bins.tsv')
for s, name in (('yeasts', 'Yeasts'), ('caenorhabditis', 'Caenorhabditis'), ('drosophila', 'Drosophila'), ('primates', 'Primates')):
    add(name, name, ROOT / 'Eukaryote_Screen' / s / 'run' / 'reports' / 'proposed_bins.tsv', eukaryote=True)

# ---------- write ----------
HEAD = PatternFill('solid', fgColor='1F3A5F'); HFONT = Font(bold=True, color='FFFFFF')
YES = PatternFill('solid', fgColor='D9EAD3'); NO = PatternFill('solid', fgColor='F4CCCC')
BAND = PatternFill('solid', fgColor='F3F6FA')
thin = Side(style='thin', color='C8D0DA')
wb = Workbook(); about = wb.active; about.title = 'About'
summary = []
for sheet, rows in groups.items():
    ws = wb.create_sheet(sheet[:31])
    nb = max(len(r) for r in rows) - 4
    hdr = ['DB_searched', 'DB_agreement', 'Genus', 'Screened set'] + [f'RaGCAn_{i+1}' for i in range(nb)]
    ws.append(hdr)
    for c in ws[1]:
        c.fill, c.font = HEAD, HFONT
        c.alignment = Alignment(horizontal='center', vertical='center')
    sets = []
    for r in rows:
        ws.append(r + [''] * (len(hdr) - len(r)))
        if not sets or sets[-1] != r[3]: sets.append(r[3])
    band = False; prev = None
    for row in ws.iter_rows(min_row=2):
        if row[3].value != prev: band = not band; prev = row[3].value
        for c in row:
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical='top', wrap_text=c.column > 4)
            if band and c.column != 2: c.fill = BAND
        row[1].fill = YES if row[1].value == 'yes' else NO
        row[1].alignment = Alignment(horizontal='center', vertical='top')
        row[2].font = Font(italic=True)
    for i, w in enumerate([13, 14, 26, 24] + [34] * nb, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"
    y = sum(r[1] == 'yes' for r in rows)
    summary.append((sheet, len(sets), len(rows), y, len(rows) - y, sum(r[0] == 'NCBI' for r in rows), nb))

src = wb.create_sheet('Sources')
src.append(['Accession', 'Organism name', 'NCBI genome page', 'Protein file downloaded'])
for c in src[1]: c.fill, c.font = HEAD, HFONT
accs = sorted({a for rows in groups.values() for r in rows for cell in r[4:] for a in cell.split(', ') if a})
for a in accs:
    n, g, ftp = info[a]
    page = f'https://www.ncbi.nlm.nih.gov/datasets/genome/{a}/'
    dl = f"{ftp}/{ftp.rsplit('/', 1)[1]}_protein.faa.gz" if ftp else ''
    src.append([a, n, page, dl])
    r = src.max_row
    src.cell(r, 3).hyperlink = page; src.cell(r, 3).style = 'Hyperlink'
    if dl: src.cell(r, 4).hyperlink = dl; src.cell(r, 4).style = 'Hyperlink'
    src.cell(r, 2).font = Font(italic=True)
for i, w in enumerate([18, 48, 62, 110], 1): src.column_dimensions[get_column_letter(i)].width = w
src.freeze_panes = 'A2'; src.auto_filter.ref = f'A1:D{src.max_row}'

about['A1'] = 'RaGCAn results workbook'; about['A1'].font = Font(bold=True, size=16, color='1F3A5F')
lines = [
 'One sheet per organism group. One row per database genus within one screened set; bins are RaGCAn at 65% AAI.',
 'DB_searched: GTDB (release R232) for prokaryote genomes that GTDB classifies; NCBI for eukaryotes and for prokaryote genomes absent from GTDB.',
 'DB_agreement: yes when all of this genus\'s genomes in the set fall in one RaGCAn bin and that bin holds no other genus; otherwise no.',
 'Screened set: the group of genomes run together (a named genus, the Enterobacteriaceae family, or a eukaryote set).',
 'RaGCAn_1, RaGCAn_2, ...: accessions of this genus\'s genomes in each bin of that set. Sources sheet links every accession to its NCBI page and the exact protein file downloaded.',
 'Note: DB_agreement is a per-genus reading and is not the same as the per-set scoring in the paper (which needs >= 3 GTDB-classified genomes).',
 '']
for i, t in enumerate(lines, 3): about.cell(i, 1, t)
r0 = 3 + len(lines)
for j, h in enumerate(['Sheet', 'Screened sets', 'Genus rows', 'DB_agreement yes', 'DB_agreement no', 'NCBI rows', 'Max bins'], 1):
    c = about.cell(r0, j, h); c.fill, c.font = HEAD, HFONT
for i, s in enumerate(summary, r0 + 1):
    for j, v in enumerate(s, 1): about.cell(i, j, v)
about.cell(r0 + len(summary) + 2, 1, f'Genomes listed in Sources: {len(accs)}')
about.column_dimensions['A'].width = 30
for col in 'BCDEFG': about.column_dimensions[col].width = 16
wb.save(OUT)
for s in summary: print(s)
print('genomes', len(accs), '->', OUT)
